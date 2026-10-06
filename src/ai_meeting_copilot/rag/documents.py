"""Loading files and splitting them into retrievable chunks."""

from __future__ import annotations

import re
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".csv", ".json", ".yaml", ".yml"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | {".pdf", ".docx"}


@dataclass(frozen=True)
class Document:
    source: str  # e.g. "resume.pdf#p2"
    text: str


@dataclass(frozen=True)
class Chunk:
    id: int  # position in the chunk list; also the Qdrant point id
    source: str
    text: str


def _expand(paths: Iterable[Path]) -> list[Path]:
    files: list[Path] = []
    for path in paths:
        if path.is_dir():
            files.extend(sorted(p for p in path.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES))
        elif path.is_file():
            files.append(path)
        else:
            raise FileNotFoundError(path)
    return files


def load_documents(paths: Iterable[Path]) -> list[Document]:
    """Read text from files/directories. PDFs become one Document per page so answers can cite pages."""
    docs: list[Document] = []
    for file in _expand(paths):
        suffix = file.suffix.lower()
        if suffix == ".pdf":
            from pypdf import PdfReader

            for page_no, page in enumerate(PdfReader(file).pages, start=1):
                text = page.extract_text() or ""
                if text.strip():
                    docs.append(Document(f"{file.name}#p{page_no}", text))
        elif suffix == ".docx":
            import docx

            text = "\n\n".join(p.text for p in docx.Document(str(file)).paragraphs)
            if text.strip():
                docs.append(Document(file.name, text))
        elif suffix in TEXT_SUFFIXES:
            text = file.read_text(encoding="utf-8", errors="replace")
            if text.strip():
                docs.append(Document(file.name, text))
        else:
            raise ValueError(f"Unsupported file type: {file.name} (supported: {', '.join(sorted(SUPPORTED_SUFFIXES))})")
    return docs


_SENTENCE_END = re.compile(r"(?<=[.!?])\s+")


def _split_units(text: str, chunk_size: int) -> list[str]:
    """Break text into pieces no longer than chunk_size: paragraphs, else sentences, else hard cuts."""
    units: list[str] = []
    for para in re.split(r"\n\s*\n", text):
        para = re.sub(r"\s+", " ", para).strip()
        if not para:
            continue
        if len(para) <= chunk_size:
            units.append(para)
            continue
        for sentence in _SENTENCE_END.split(para):
            while len(sentence) > chunk_size:
                units.append(sentence[:chunk_size])
                sentence = sentence[chunk_size:]
            if sentence:
                units.append(sentence)
    return units


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 150) -> list[str]:
    """Greedy packing of paragraphs/sentences into ~chunk_size characters.

    Each new chunk starts with the trailing units of the previous one (up to `overlap`
    characters), so a fact that straddles a boundary is still retrievable whole.
    """
    if overlap >= chunk_size:
        raise ValueError("overlap must be smaller than chunk_size")
    chunks: list[str] = []
    current: list[str] = []
    length = 0

    for unit in _split_units(text, chunk_size):
        if current and length + 1 + len(unit) > chunk_size:
            chunks.append(" ".join(current))
            tail: list[str] = []
            tail_len = 0
            for prev in reversed(current):
                if tail_len + len(prev) + 1 > overlap:
                    break
                tail.insert(0, prev)
                tail_len += len(prev) + 1
            # Never let the carried-over tail plus the new unit overflow the chunk.
            while tail and tail_len + len(unit) > chunk_size:
                tail_len -= len(tail.pop(0)) + 1
            current, length = tail, tail_len
        current.append(unit)
        length += len(unit) + (1 if length else 0)

    if current:
        chunks.append(" ".join(current))
    return chunks


def chunk_documents(docs: Iterable[Document], chunk_size: int = 800, overlap: int = 150) -> list[Chunk]:
    chunks: list[Chunk] = []
    for doc in docs:
        for piece in chunk_text(doc.text, chunk_size, overlap):
            chunks.append(Chunk(id=len(chunks), source=doc.source, text=piece))
    return chunks
